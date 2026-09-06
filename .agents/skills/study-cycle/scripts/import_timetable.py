#!/usr/bin/env python3
"""Import an Excel timetable while preserving stable study-session identifiers."""

from __future__ import annotations

import argparse
import base64
import json
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timezone, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SEOUL = timezone(timedelta(hours=9))
EXAM_TITLES = {"시험", "중간고사", "기말고사", "평가"}


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalized(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).lower()
    return re.sub(r"[^0-9a-z가-힣]+", "", value)


def base_title(value: str) -> str:
    value = re.sub(r"\s*[\[(]?\d+[\])]?[\s]*$", "", value.strip())
    return normalized(value)


def to_datetime(day: Any, clock: Any) -> datetime:
    if isinstance(day, datetime):
        day = day.date()
    if not isinstance(day, date):
        raise ValueError(f"invalid date cell: {day!r}")
    if isinstance(clock, datetime):
        clock = clock.time()
    if not isinstance(clock, time):
        raise ValueError(f"invalid time cell: {clock!r}")
    return datetime.combine(day, clock, tzinfo=SEOUL)


def iso(value: datetime) -> str:
    return value.isoformat(timespec="minutes")


@dataclass
class WorkbookRow:
    row: int
    week: int
    period: int
    scheduled_at: datetime
    professor: str
    title: str
    room: str
    online: bool
    change_note: str | None


def read_rows(path: Path, sheet_name: str, course_name: str) -> tuple[list[WorkbookRow], dict[str, str] | None]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"sheet {sheet_name!r} not found; available: {workbook.sheetnames}")
    sheet = workbook[sheet_name]
    rows: list[WorkbookRow] = []
    exam: dict[str, str] | None = None
    exam_keys = {normalized(item) for item in EXAM_TITLES}
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, max_col=10, values_only=True), start=2):
        week, day, period, clock, course, professor, title, room, lms, change_note = values
        if text(course) != course_name:
            continue
        scheduled_at = to_datetime(day, clock)
        title_text = text(title)
        if normalized(title_text) in exam_keys:
            exam = {"scheduledAt": iso(scheduled_at), "room": text(room)}
            continue
        rows.append(
            WorkbookRow(
                row=row_number,
                week=int(week),
                period=int(period),
                scheduled_at=scheduled_at,
                professor=text(professor),
                title=title_text,
                room=text(room),
                online=text(lms).upper() == "O",
                change_note=text(change_note) or None,
            )
        )
    rows.sort(key=lambda item: (item.scheduled_at, item.period, item.row))
    return rows, exam


def load_json(path: Path | None, default: Any) -> Any:
    if path is None or not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def build_session(row: WorkbookRow, sequence: int, slug: str, previous: dict[str, Any] | None) -> dict[str, Any]:
    scheduled_at = iso(row.scheduled_at)
    return {
        "id": previous["id"] if previous else f"{slug}-{sequence:03d}",
        "sequence": sequence,
        "week": row.week,
        "period": row.period,
        "scheduledAt": scheduled_at,
        "originalScheduledAt": previous.get("originalScheduledAt", previous.get("scheduledAt")) if previous else scheduled_at,
        "title": row.title,
        "professor": row.professor,
        "room": row.room,
        "online": row.online,
        "changeNote": row.change_note,
        "status": "scheduled",
        "sourceGroupId": previous.get("sourceGroupId") if previous else None,
        "notebookUrl": previous.get("notebookUrl") if previous else None,
    }


def match_rows(rows: list[WorkbookRow], previous: dict[str, Any] | None, slug: str) -> list[dict[str, Any]]:
    if not previous:
        return [build_session(row, index, slug, None) for index, row in enumerate(rows, start=1)]

    old_sessions = list(previous.get("sessions", []))
    unused = {item["id"]: item for item in old_sessions}
    max_sequence = max((int(item["sequence"]) for item in old_sessions), default=0)
    matched: list[dict[str, Any]] = []
    exact_groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    base_groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for item in old_sessions:
        exact_groups[(normalized(item["title"]), normalized(item.get("professor", "")))].append(item)
        base_groups[(base_title(item["title"]), normalized(item.get("professor", "")))].append(item)

    for row in rows:
        key = (normalized(row.title), normalized(row.professor))
        candidates = [item for item in exact_groups[key] if item["id"] in unused]
        if not candidates:
            base_key = (base_title(row.title), normalized(row.professor))
            candidates = [item for item in base_groups[base_key] if item["id"] in unused]

        chosen = None
        review = False
        if candidates:
            distances = sorted(
                (
                    abs((datetime.fromisoformat(item["scheduledAt"]) - row.scheduled_at).total_seconds()),
                    int(item["sequence"]),
                    item,
                )
                for item in candidates
            )
            chosen = distances[0][2]
            review = len(distances) > 1 and distances[0][0] == distances[1][0]

        if chosen:
            unused.pop(chosen["id"], None)
            session = build_session(row, int(chosen["sequence"]), slug, chosen)
            if review:
                session["status"] = "review"
            matched.append(session)
        else:
            max_sequence += 1
            session = build_session(row, max_sequence, slug, None)
            session["status"] = "review"
            matched.append(session)

    for missing in unused.values():
        retained = dict(missing)
        retained["status"] = "cancelled"
        matched.append(retained)

    matched.sort(key=lambda item: int(item["sequence"]))
    return matched


def apply_materials(sessions: list[dict[str, Any]], raw_materials: Any) -> list[dict[str, Any]]:
    groups = raw_materials.get("sourceGroups", raw_materials if isinstance(raw_materials, list) else [])
    sequence_map = {int(item["sequence"]): item for item in sessions}
    output = []
    for index, raw in enumerate(groups, start=1):
        sequences = [int(item) for item in raw["sessionSequences"]]
        group_id = raw.get("id") or f"source-{index:03d}"
        for sequence in sequences:
            if sequence not in sequence_map:
                raise ValueError(f"material group {group_id} references unknown session {sequence}")
            sequence_map[sequence]["sourceGroupId"] = group_id
            if raw.get("notebookUrl"):
                sequence_map[sequence]["notebookUrl"] = raw["notebookUrl"]
        first, last = min(sequences), max(sequences)
        range_label = str(first) if first == last else f"{first}~{last}"
        title = raw.get("title") or sequence_map[first]["title"]
        professor = raw.get("professor") or sequence_map[first]["professor"]
        output.append(
            {
                "id": group_id,
                "sessionSequences": sequences,
                "title": title,
                "professor": professor,
                "fileName": raw.get("fileName"),
                "driveUrl": raw.get("driveUrl"),
                "notebookName": raw.get("notebookName") or f"{range_label}. {title}({professor})",
                "notebookUrl": raw.get("notebookUrl"),
                "materialStatus": raw.get("materialStatus", "ready" if raw.get("fileName") else "waiting"),
                "warmup": raw.get("warmup"),
                "recall": raw.get("recall"),
                "noteGuide": raw.get("noteGuide"),
            }
        )
    return output


def write_env(path: Path, manifest: dict[str, Any]) -> None:
    compact = json.dumps(manifest, ensure_ascii=False, separators=(",", ":"))
    encoded = base64.b64encode(compact.encode("utf-8")).decode("ascii")
    path.write_text(f"COURSE_MANIFEST_BASE64={encoded}\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--course", required=True)
    parser.add_argument("--slug", required=True)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--previous", type=Path)
    parser.add_argument("--materials", type=Path)
    parser.add_argument("--env-file", type=Path)
    parser.add_argument("--cycle-label", default="1 BLOCK")
    parser.add_argument("--drive-folder-url")
    args = parser.parse_args()

    rows, exam = read_rows(args.workbook, args.sheet, args.course)
    previous = load_json(args.previous, None)
    sessions = match_rows(rows, previous, args.slug)
    source_groups = apply_materials(sessions, load_json(args.materials, []))
    manifest = {
        "schemaVersion": 1,
        "slug": args.slug,
        "name": args.course,
        "cycleLabel": args.cycle_label,
        "timezone": "Asia/Seoul",
        "sourceVersion": args.workbook.stem,
        "driveFolderUrl": args.drive_folder_url,
        "sessions": sessions,
        "sourceGroups": source_groups,
        "exam": exam,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.env_file:
        write_env(args.env_file, manifest)

    review = [item["sequence"] for item in sessions if item["status"] == "review"]
    cancelled = [item["sequence"] for item in sessions if item["status"] == "cancelled"]
    print(json.dumps({"sessions": len(sessions), "sourceGroups": len(source_groups), "review": review, "cancelled": cancelled, "exam": bool(exam)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
